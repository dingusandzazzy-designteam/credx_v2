# -*- coding: utf-8 -*-
# Generate the client-ready CredX SEO/AEO Strategy v2 .docx from the markdown master.
# Source of truth: credx-website/seo/CredX-SEO-AEO-Strategy-v2.md
import os, re
from docx import Document
from docx.shared import Pt, RGBColor

SRC = r"credx-website/seo/CredX-SEO-AEO-Strategy-v2.md"
OUT = r"credx-website/documents/02.seo/CredX SEO_AEO Strategy v2.docx"

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def add_runs(p, text):
    """Render inline **bold** and `code` (backticks stripped) into paragraph p."""
    text = text.replace('`', '')
    # split on bold markers, keeping delimiters via regex
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        else:
            p.add_run(part)

def flush_table(rows):
    # rows: list of list-of-cell-strings; row 0 = header
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = 'Light Grid Accent 1'
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            val = row[ci] if ci < len(row) else ''
            cell = t.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            add_runs(p, val)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph()

with open(SRC, encoding='utf-8') as f:
    lines = f.read().split('\n')

i = 0
n = len(lines)
while i < n:
    line = lines[i].rstrip('\n')
    stripped = line.strip()

    # blank
    if not stripped:
        i += 1
        continue

    # horizontal rule
    if re.fullmatch(r'-{3,}', stripped):
        i += 1
        continue

    # table block
    if stripped.startswith('|'):
        block = []
        while i < n and lines[i].strip().startswith('|'):
            block.append(lines[i].strip())
            i += 1
        parsed = []
        for b in block:
            # skip separator row |---|---|
            if re.fullmatch(r'\|[\s:|-]+\|', b):
                continue
            cells = [c.strip() for c in b.strip('|').split('|')]
            parsed.append(cells)
        if parsed:
            flush_table(parsed)
        continue

    # headings
    m = re.match(r'(#{1,6})\s+(.*)', stripped)
    if m:
        hashes, txt = m.group(1), m.group(2)
        txt = txt.replace('`', '')
        level = len(hashes) - 1  # '#' -> 0 (title), '##' -> 1, ...
        if level > 4:
            level = 4
        doc.add_heading(txt, level=level)
        i += 1
        continue

    # blockquote -> italic note paragraph(s)
    if stripped.startswith('>'):
        qtext = re.sub(r'^>\s?', '', stripped)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        # render inline bold but force italic on the plain runs
        qtext = qtext.replace('`', '')
        parts = re.split(r'(\*\*.+?\*\*)', qtext)
        for part in parts:
            if not part:
                continue
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2]); r.bold = True; r.italic = True
            else:
                r = p.add_run(part); r.italic = True
        i += 1
        continue

    # bullet
    if re.match(r'[-*]\s+', stripped):
        btext = re.sub(r'^[-*]\s+', '', stripped)
        p = doc.add_paragraph(style='List Bullet')
        add_runs(p, btext)
        i += 1
        continue

    # ordered list item
    mo = re.match(r'(\d+)\.\s+(.*)', stripped)
    if mo:
        p = doc.add_paragraph(style='List Number')
        add_runs(p, mo.group(2))
        i += 1
        continue

    # default paragraph
    p = doc.add_paragraph()
    add_runs(p, stripped)
    i += 1

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print('wrote', OUT, os.path.getsize(OUT)//1024, 'KB')
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
