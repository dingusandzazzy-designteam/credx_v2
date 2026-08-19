# -*- coding: utf-8 -*-
"""
Regenerate the three Batch-1 review .docx (/about, /how-it-works, /evp)
FROM THE SERVED HTML, so a review copy can never drift from the shipped page.

Why from HTML and not from copy/*.md: the md files carry internal decision
notes, superseded strings and open questions. A review doc must carry only
what a reader actually sees. The HTML is the single unambiguous record of that.

Counts in the metadata table are MEASURED here, not copied from any brief —
the briefs were wrong on three values (corrected 2026-08-12).

Usage:  python scripts/gen_batch1_docx.py
Output: credx-website/copy/docx/*.docx   (gitignored, overwritten in place)
"""
import html
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
SITE = os.path.join(ROOT, 'credx-website')
OUT = os.path.join(SITE, 'copy', 'docx')

# Filenames are kept BYTE-IDENTICAL to the 2026-08-10 exports on purpose, so the
# 1:1 mapping to the three existing Google Docs stays obvious when pasting over
# them. Note the /evp filename still says "Partner" — the page itself now says
# "Platform" (C2). Renaming the file is a separate call, not a silent side effect.
PAGES = [
    {
        'slug': 'about',
        'file': 'CredX — The Movement (about).docx',
        'title': 'The Movement',
        'meta': '/about  ·  nav #1  ·  indexed  ·  sitemap priority 0.7',
        'context': 'Primary keyword: the data economy (50/mo, Canada). Schema: Organization + BreadcrumbList. '
                   'This is a brand and narrative page, not an acquisition page.',
    },
    {
        'slug': 'how-it-works',
        'file': 'CredX — How It Works.docx',
        'title': 'How It Works',
        'meta': '/how-it-works  ·  nav #2  ·  indexed  ·  sitemap priority 0.9',
        'context': 'Primary keyword: embedded lending. Schema: HowTo + FAQPage + BreadcrumbList. '
                   'Informational page — the Canadian SERP for this query is 100% editorial.',
        # Same classes, different roles per page: on /about `hero__sub` is a narrative
        # subheadline and `hero__movement` is a genuine movement line; here the sub IS
        # the definition answer engines quote, and the movement slot is a lead-in.
        'labels': {
            'hero__sub': 'AEO DEFINITION — the answer engines quote this',
            'hero__movement': 'LEAD-IN LINE',
        },
    },
    {
        'slug': 'evp',
        'file': 'CredX — Embedded Value Partner (EVP).docx',
        'title': 'Embedded Value Platform',
        'meta': '/evp  ·  not in nav (reachable from /about §3 and the sitemap)  ·  indexed  ·  sitemap priority 0.8',
        'context': 'AEO / definition page for a CredX-coined category. Schema: Organization + Service + BreadcrumbList. '
                   'The §1 definition snippet is the definition of record for answer engines.',
        'labels': {
            'hero__sub': 'AEO DEFINITION OF RECORD — highest-consequence string on the site',
            'hero__movement': 'SUB-LINE',
        },
    },
]

# Section name -> the first BODY inside it is really the AEO snippet, not prose.
# /about carries its definition in a normal prose paragraph, so class alone cannot
# tell them apart; position can.
FIRST_BODY_IS_AEO = {'The data economy': 'AEO DEFINITION — the answer engines quote this'}

# class -> label. Leaf-ish only, so nothing is captured twice.
LABELS = {
    'eyebrow': 'EYEBROW',
    'hero__h1': 'H1',
    'hero__sub': 'SUBHEADLINE',
    'hero__movement': 'MOVEMENT LINE',
    'prose-band__h2': 'H2',
    'control__h2': 'H2',
    'final-cta__h2': 'HEADLINE',
    'control__intro': 'INTRO',
    'control-card__title': 'CARD TITLE',
    'control-card__body': 'CARD BODY',
    'gain__title': 'COMPONENT',
    'gain__body': 'COMPONENT BODY',
    'prose-band__body': 'BODY',
    'final-cta__body': 'BODY',
    'consent-note': 'CONSENT LINE',
    'control__whitelabel-head': 'CALLOUT',
    'control__whitelabel': 'CALLOUT BODY',
    'gains__benchmark': 'ASIDE',
    'section-aside': 'ASIDE',
    'prose-band__aside': 'ASIDE',
    'step__eyebrow': 'STEP',
    'step__h3': 'STEP TITLE',
    'step__body': 'STEP BODY',
    'faq__question': 'Q',
    'faq__answer': 'A',
    'btn-primary': 'PRIMARY CTA',
    'btn-secondary': 'SECONDARY CTA',
    'sub-link': 'CROSS-LINK',
    'trust__lead': 'TRUST',
    'trust__body': 'TRUST BODY',
    'trust-note': 'TRUST NOTE',
    'stat__number': 'STAT',
    'stat__label': 'STAT CAPTION',
}
BARE = {'h1': 'H1', 'h2': 'H2', 'h3': 'H3'}
SKIP_ANCESTOR = ('nav', 'footer')  # site chrome is not page copy


def strip_tags(s):
    """Inline tags are removed WITHOUT inserting whitespace — that is how a browser
    renders them, and it is what the SEO budget must be measured against.
    Replacing tags with a space instead added a phantom character to any string with
    inline markup (`What Is <em>Embedded Lending</em>?` measured 26 rather than 25),
    which silently inflated every count that mattered here."""
    s = re.sub(r'(?s)<[^>]+>', '', s)
    return re.sub(r'\s+', ' ', html.unescape(s)).strip()


class Extract(HTMLParser):
    """Emits (label, text) in document order for one <section>."""

    def __init__(self, overrides=None):
        super().__init__(convert_charrefs=True)
        self.out = []
        self.cap = None      # (label, depth, tag, buf, href)
        self.chrome = 0
        self.labels = dict(LABELS)
        if overrides:
            self.labels.update(overrides)

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in SKIP_ANCESTOR:
            self.chrome += 1
            return
        if self.chrome:
            return

        if self.cap:
            if tag == self.cap[2]:
                self.cap[1][0] += 1
            return

        if tag == 'img':
            alt = (a.get('alt') or '').strip()
            if alt:
                self.out.append(('IMAGE ALT', alt))
            return

        classes = (a.get('class') or '').split()
        label = None
        for c in classes:
            if c in self.labels:
                label = self.labels[c]
                break
        if label is None and 'media-ph' in classes:
            self.out.append(('IMAGE SLOT', '— placeholder, no image wired —'))
            return
        if label is None and tag in BARE and not classes:
            label = BARE[tag]
        if label is None:
            return

        if tag == 'a':
            href = a.get('href', '')
            self.cap = (label, [0], tag, [], href)
        else:
            self.cap = (label, [0], tag, [], None)

    def handle_endtag(self, tag):
        if tag in SKIP_ANCESTOR:
            self.chrome = max(0, self.chrome - 1)
            return
        if self.chrome or not self.cap:
            return
        if tag != self.cap[2]:
            return
        if self.cap[1][0] > 0:
            self.cap[1][0] -= 1
            return
        label, _, _, buf, href = self.cap
        text = re.sub(r'\s+', ' ', ''.join(buf)).strip()
        self.cap = None
        if not text:
            return
        if href and not href.startswith('#'):
            text = f'{text}  →  {href}'
        self.out.append((label, text))

    def handle_data(self, data):
        if self.cap and not self.chrome:
            self.cap[3].append(data)


def sections(page_html):
    """(human name, inner html) per <section>, in document order."""
    h = re.sub(r'(?s)<!--.*?-->', '', page_html)
    res = []
    for m in re.finditer(r'(?s)<section([^>]*)>(.*?)</section>', h):
        attrs, inner = m.group(1), m.group(2)
        lbl = re.search(r'aria-label="([^"]+)"', attrs)
        res.append((lbl.group(1) if lbl else 'Section', inner))
    return res


def measure(page_html):
    t = re.search(r'<title>(.*?)</title>', page_html, re.S)
    d = re.search(r'<meta name="description" content="([^"]*)"', page_html)
    h1 = re.search(r'(?s)<h1[^>]*>(.*?)</h1>', page_html)
    return {
        'Title': (strip_tags(t.group(1)) if t else '', 60),
        'Meta description': (html.unescape(d.group(1)) if d else '', 160),
        'H1': (strip_tags(h1.group(1)) if h1 else '', 70),
    }


def build(page):
    src = os.path.join(SITE, page['slug'], 'index.html')
    raw = open(src, encoding='utf-8').read()

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    doc.add_heading(page['title'], level=0)

    p = doc.add_paragraph()
    p.add_run(page['meta']).bold = True

    p = doc.add_paragraph()
    r = p.add_run('DRAFT 2026-08-12 — regenerated from the served HTML, not human-reviewed, '
                  'not client-reviewed. Working file: copy/%s.md' % page['slug'])
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    p = doc.add_paragraph()
    p.add_run(page['context']).italic = True

    doc.add_heading('Page metadata', level=1)
    m = measure(raw)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0].cells
    for i, t in enumerate(('Field', 'Value', 'Measured / budget')):
        hdr[i].text = ''
        hdr[i].paragraphs[0].add_run(t).bold = True
    for field, (val, budget) in m.items():
        cells = tbl.add_row().cells
        cells[0].text = field
        cells[1].text = val
        cells[2].text = f'{len(val)} / {budget}'
        if field == 'Meta description' and len(val) >= budget:
            cells[2].paragraphs[0].runs[0].bold = True

    n = 0
    for name, inner in sections(raw):
        ex = Extract(page.get('labels'))
        ex.feed(inner)
        fields = ex.out
        if not fields:
            continue

        # /about's definition paragraph is an ordinary prose paragraph, so only its
        # position identifies it. Relabel the first BODY in the named section.
        if name in FIRST_BODY_IS_AEO:
            for idx, (lb, tx) in enumerate(fields):
                if lb == 'BODY':
                    fields[idx] = (FIRST_BODY_IS_AEO[name], tx)
                    break

        # Image references go last, as in the 2026-08-10 exports. In markup the <img>
        # often precedes the copy, and a review doc that opens on alt text reads wrong.
        media = [f for f in fields if f[0] in ('IMAGE ALT', 'IMAGE SLOT')]
        fields = [f for f in fields if f[0] not in ('IMAGE ALT', 'IMAGE SLOT')] + media

        n += 1
        doc.add_heading(f'{n} · {name}', level=1)
        i = 0
        while i < len(fields):
            label, text = fields[i]
            # pair STAT + its caption onto one line
            if label == 'STAT' and i + 1 < len(fields) and fields[i + 1][0] == 'STAT CAPTION':
                text = f'{text} — {fields[i + 1][1]}'
                i += 1
            lp = doc.add_paragraph()
            lr = lp.add_run(label)
            lr.bold = True
            lr.font.size = Pt(8)
            lp.paragraph_format.space_after = Pt(0)
            vp = doc.add_paragraph(text)
            vp.paragraph_format.space_before = Pt(0)
            if label in ('Q',):
                vp.runs[0].bold = True
            i += 1

    doc.add_page_break()
    doc.add_heading('How this document was produced', level=1)
    for line in (
        'Generated by scripts/gen_batch1_docx.py directly from credx-website/%s/index.html.' % page['slug'],
        'Every string below the metadata table is the rendered page text — site navigation and footer excluded.',
        'Character counts in the metadata table are measured from this file, not taken from the SEO brief. '
        'The briefs carried three wrong values, corrected 2026-08-12.',
        'Internal decision notes, superseded strings and open questions are deliberately NOT included. '
        'Those live in copy/%s.md.' % page['slug'],
    ):
        b = doc.add_paragraph(line, style='List Bullet')
        b.runs[0].font.size = Pt(9)
        b.runs[0].italic = True

    os.makedirs(OUT, exist_ok=True)
    dest = os.path.join(OUT, page['file'])
    doc.save(dest)
    return dest, n, m


if __name__ == '__main__':
    for pg in PAGES:
        dest, n, m = build(pg)
        print(f"OK  {os.path.basename(dest)}  ({n} sections)")
        for f, (v, b) in m.items():
            flag = '  <-- AT/OVER CEILING' if f == 'Meta description' and len(v) >= b else ''
            print(f"      {f}: {len(v)} / {b}{flag}")
